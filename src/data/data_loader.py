"""
Унифицированный загрузчик данных Санкт-Петербургского метрополитена (Линия 1).
Читает 15-минутные помесячные файлы (февраль, май, июль, сентябрь 2026),
8-месячный почасовой датасет, графики оборота составов и показатели МО-1.
"""

from typing import Optional, Dict, Any, List, Tuple
import os
import re
import datetime
import openpyxl
import pandas as pd
import numpy as np

import sys
# Обеспечиваем импорт из корневого пакета configs
sys.path.append(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
from configs.stations import get_station_info, STATIONS_LINE_1
from configs.metro_config import METRO_CONFIG
from src.data.operating_hours import filter_operating_hours


class MetroDataLoader:
    """
    Класс загрузки и первичной обработки телеметрии пассажиропотока Линии 1.
    """
    
    def __init__(self, data_root: Optional[str] = None):
        if data_root is None:
            base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
            self.data_root = os.path.join(base_dir, "data")
        else:
            self.data_root = data_root
            
        self.flow_dir = os.path.join(self.data_root, "1.Пассажиропоток+ ГД")
        self.mo1_dir = os.path.join(self.data_root, "2.ОСНОВНЫЕ ПОКАЗАТЕЛИ выполнения ГД")
        self.specs_dir = os.path.join(self.data_root, "3.Характеристики состава и Линии")
        
        self.cache_parquet = os.path.join(self.data_root, "processed_15min_line1.parquet")

    def _parse_operational_sheet(self, ws, sname: str) -> List[Dict[str, Any]]:
        """Парсит один суточный лист 15-минутного пассажиропотока через iter_rows."""
        # Читаем все строки сразу
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 4:
            return []
            
        # 1. Извлекаем дату
        title_val = rows[0][0] if rows[0] and rows[0][0] is not None else ""
        m = re.search(r'(\d{2})(\d{2})(\d{4})', str(title_val) + " " + sname)
        if not m:
            return []
            
        day, month, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
        op_date = datetime.date(year, month, day)
        
        # 2. Извлекаем временные слоты (ряд 1, столбцы 2..97)
        time_row = rows[1]
        time_slots = []
        for c in range(2, 98):
            cell_val = time_row[c] if c < len(time_row) else None
            if isinstance(cell_val, datetime.time):
                time_slots.append(cell_val)
            elif isinstance(cell_val, str) and ":" in cell_val:
                h, mi = map(int, cell_val.strip().split(":")[:2])
                time_slots.append(datetime.time(h, mi))
            else:
                idx = c - 2
                total_min = (3 * 60 + idx * 15) % (24 * 60)
                time_slots.append(datetime.time(total_min // 60, total_min % 60))
                
        # 3. Читаем вестибюли (ряды 3..26)
        records = []
        for r in range(3, min(27, len(rows))):
            row_data = rows[r]
            if not row_data or row_data[0] is None:
                continue
            vest_raw = str(row_data[0]).strip()
            if not vest_raw:
                continue
                
            try:
                st_info = get_station_info(vest_raw)
            except KeyError:
                continue
                
            for idx, t_obj in enumerate(time_slots):
                col_idx = 2 + idx
                val = row_data[col_idx] if col_idx < len(row_data) else 0.0
                p_count = float(val) if isinstance(val, (int, float)) else 0.0
                
                # Обработка перехода через полночь (00:00 - 02:45)
                if t_obj.hour < 3:
                    full_dt = datetime.datetime.combine(op_date + datetime.timedelta(days=1), t_obj)
                else:
                    full_dt = datetime.datetime.combine(op_date, t_obj)
                    
                records.append({
                    "datetime": full_dt,
                    "operational_date": op_date,
                    "time_slot": t_obj.strftime("%H:%M"),
                    "interval_index": idx,
                    "raw_vestibule": vest_raw,
                    "vestibule": st_info["vestibule"],
                    "station_code": st_info["station_code"],
                    "station_name": st_info["station_name"],
                    "station_order": st_info["order"],
                    "station_type": st_info["type"],
                    "lat": st_info["lat"],
                    "lon": st_info["lon"],
                    "has_turnaround": int(st_info["has_turnaround"]),
                    "passengers": p_count
                })
                
        return records

    def load_15min_dataset(self, force_reload: bool = False) -> pd.DataFrame:
        """
        Загружает и объединяет все 4 месяца (февраль, май, июль, сентябрь 2026)
        15-минутного пассажиропотока по вестибюлям.

        Возвращаются только пассажирские слоты 05:30-00:30 (see operating_hours.py):
        ночные проходы сотрудников (00:30-05:15) не входят в пассажиропоток.
        """
        if os.path.exists(self.cache_parquet) and not force_reload:
            try:
                return filter_operating_hours(pd.read_parquet(self.cache_parquet))
            except Exception:
                pass
                
        files = [
            "Входные пассажиропотоки Линия 1 фев 2026 по 15-мин.xlsx",
            "Входные пассажиропотоки Линия 1 май 2026 по 15-мин.xlsx",
            "Входные пассажиропотоки Линия 1 июль 2026 по 15-мин.xlsx",
            "Входные пассажиропотоки Линия 1 сен 2026 по 15-мин.xlsx"
        ]
        
        all_records = []
        for fname in files:
            fpath = os.path.join(self.flow_dir, fname)
            if not os.path.exists(fpath):
                continue
            wb = openpyxl.load_workbook(fpath, data_only=True, read_only=True)
            for sname in wb.sheetnames:
                if "Лист" in sname:
                    continue
                ws = wb[sname]
                sheet_records = self._parse_operational_sheet(ws, sname)
                all_records.extend(sheet_records)
            wb.close()
            
        df = pd.DataFrame(all_records)
        df["datetime"] = pd.to_datetime(df["datetime"])
        df = df.sort_values(["datetime", "station_order", "vestibule"]).reset_index(drop=True)
        
        # Сохранение в кэш
        os.makedirs(os.path.dirname(self.cache_parquet), exist_ok=True)
        df.to_parquet(self.cache_parquet, index=False)  # кэш хранит сырые данные целиком
        return filter_operating_hours(df)

    def get_station_aggregated_flow(self, df_15min: Optional[pd.DataFrame] = None) -> pd.DataFrame:
        """Агрегирует пассажиропоток по станциям (суммируя вестибюли)."""
        if df_15min is None:
            df_15min = self.load_15min_dataset()
            
        group_cols = [
            "datetime", "operational_date", "time_slot", "interval_index",
            "station_code", "station_name", "station_order", "station_type",
            "lat", "lon", "has_turnaround"
        ]
        
        df_station = df_15min.groupby(group_cols, as_index=False)["passengers"].sum()
        return df_station.sort_values(["datetime", "station_order"]).reset_index(drop=True)

    def get_line_aggregated_flow(self, df_15min: Optional[pd.DataFrame] = None) -> pd.DataFrame:
        """Агрегирует пассажиропоток по всей Линии 1 в целом на каждом 15-минутном интервале."""
        if df_15min is None:
            df_15min = self.load_15min_dataset()
            
        group_cols = ["datetime", "operational_date", "time_slot", "interval_index"]
        df_line = df_15min.groupby(group_cols, as_index=False)["passengers"].sum()
        return df_line.sort_values("datetime").reset_index(drop=True)

    def get_planned_schedule_profiles(self) -> Dict[str, pd.DataFrame]:
        """
        Возвращает официальные почасовые графики парности и числа составов на Линии 1
        (извлечено из официальных таблиц оборота составов электродепо ТЧ-1 'Автово').
        """
        # Рабочий график (сентябрь - максимальная интенсивность)
        workday_sep = pd.DataFrame({
            "hour": list(range(5, 24)) + [0, 1],
            "trains_on_line": [12, 33, 50, 53, 43, 36, 33, 30, 30, 33, 40, 46, 50, 50, 46, 33, 25, 23, 18, 12, 0],
            "pairs_per_hour": [7, 20, 30, 32, 26, 22, 20, 18, 18, 20, 24, 28, 30, 30, 28, 20, 15, 14, 11, 7, 0],
        })
        workday_sep["headway_sec"] = np.where(workday_sep["pairs_per_hour"] > 0, 3600.0 / workday_sep["pairs_per_hour"], 0.0)
        
        # Выходной график
        weekend_sched = pd.DataFrame({
            "hour": list(range(5, 24)) + [0, 1],
            "trains_on_line": [11, 26, 31, 35, 33, 33, 33, 31, 31, 31, 33, 33, 33, 31, 30, 26, 25, 21, 18, 12, 0],
            "pairs_per_hour": [7, 16, 19, 21, 20, 20, 20, 19, 19, 19, 20, 20, 20, 19, 18, 16, 15, 13, 11, 7, 0],
        })
        weekend_sched["headway_sec"] = np.where(weekend_sched["pairs_per_hour"] > 0, 3600.0 / weekend_sched["pairs_per_hour"], 0.0)
        
        # Летний рабочий график (июнь - август)
        workday_summer = pd.DataFrame({
            "hour": list(range(5, 24)) + [0, 1],
            "trains_on_line": [7, 33, 46, 49, 41, 35, 31, 30, 30, 33, 36, 41, 45, 48, 41, 33, 25, 23, 15, 12, 0],
            "pairs_per_hour": [5, 20, 28, 30, 25, 21, 19, 18, 18, 20, 22, 25, 27, 29, 25, 20, 15, 14, 9, 7, 0],
        })
        workday_summer["headway_sec"] = np.where(workday_summer["pairs_per_hour"] > 0, 3600.0 / workday_summer["pairs_per_hour"], 0.0)
        
        return {
            "workday_autumn": workday_sep,
            "workday_summer": workday_summer,
            "weekend": weekend_sched
        }

    def load_hourly_annual_dataset(self, force_reload: bool = False) -> pd.DataFrame:
        """Загружает 8-месячный почасовой датасет (Пассажиропоток 2026 Линия 1.xlsx)."""
        cache_hourly_path = os.path.join(self.data_root, "processed_hourly_annual.parquet")
        if os.path.exists(cache_hourly_path) and not force_reload:
            try:
                return pd.read_parquet(cache_hourly_path)
            except Exception:
                pass
                
        big_path = os.path.join(self.flow_dir, "Пассажиропоток 2026 Линия 1.xlsx")
        df_raw = pd.read_excel(big_path, skiprows=6)
        df_raw.columns = ["date_hour_str", "raw_vestibule", "passengers"]
        # Фильтруем итоговую строчку внизу
        df_raw = df_raw.dropna(subset=["date_hour_str", "raw_vestibule"]).copy()
        df_raw = df_raw[~df_raw["date_hour_str"].str.contains("Общее", na=False)]
        
        # Парсим дату-час (DD.MM.YYYY HH)
        df_raw["datetime"] = pd.to_datetime(df_raw["date_hour_str"], format="%d.%m.%Y %H", errors="coerce")
        df_raw = df_raw.dropna(subset=["datetime"])
        
        # Привязываем станции
        def map_st(v):
            try:
                info = get_station_info(str(v))
                return pd.Series([info["station_code"], info["station_name"], info["order"], info["type"]])
            except KeyError:
                return pd.Series([None, None, None, None])
                
        cols = ["station_code", "station_name", "station_order", "station_type"]
        df_raw[cols] = df_raw["raw_vestibule"].apply(map_st)
        res = df_raw.sort_values(["datetime", "station_order"]).reset_index(drop=True)
        res.to_parquet(cache_hourly_path, index=False)
        return res
